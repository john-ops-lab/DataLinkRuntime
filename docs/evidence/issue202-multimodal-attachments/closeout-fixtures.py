"""Generate the independent A4 source and frozen oracle; never call DLR.

Run from backend with the optional fixture-only Pillow dependency:
uv run --with Pillow python PATH_TO_THIS_SCRIPT OUTPUT
This does not add a runtime dependency or call a model.
"""

import hashlib
import io
import json
import sys
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.drawing.image import Image
from PIL import Image as Pixels
from PIL import ImageDraw


def generate(output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    target = output / "composite-a4.xlsx"
    oracle_path = output / "composite-a4-oracle.json"
    if target.exists() or oracle_path.exists():
        raise SystemExit("Refuse to replace frozen source/oracle")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Facts"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Quarterly closeout"
    rows = [
        ["Quarter", "Revenue", "Growth", "Date", "Zero", "Flag"],
        ["Q1", 10, 0.125, datetime(2026, 1, 15), 0, False],
        ["Q2", 20, 0.25, datetime(2026, 4, 15), None, None],
        ["Q3", 30, 0.375, datetime(2026, 7, 15), None, None],
        ["Total", "=SUM(B3:B5)", None, "=SUM(B3:B5)*2", None, None],
    ]
    for row in rows:
        sheet.append(row)
    for row in range(3, 6):
        sheet[f"B{row}"].number_format = '"$"#,##0.00'
        sheet[f"C{row}"].number_format = "0.0%"
        sheet[f"D{row}"].number_format = "yyyy-mm-dd"
    chart = BarChart()
    chart.title = "Quarterly revenue"
    chart.add_data(Reference(sheet, min_col=2, min_row=2, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=5))
    chart.style = 10
    for index, color in enumerate(["2563EB", "F59E0B", "16A34A"]):
        point = DataPoint(idx=index)
        point.graphicalProperties.solidFill = color
        chart.series[0].data_points.append(point)
    sheet.add_chart(chart, "A9")
    pixels = Pixels.new("RGB", (360, 96), "white")
    draw = ImageDraw.Draw(pixels)
    draw.rectangle((5, 5, 354, 90), fill="#f59e0b")
    draw.text((24, 36), "EMBEDDED-ORANGE-7C19", fill="black")
    image = io.BytesIO()
    pixels.save(image, format="PNG")
    image.seek(0)
    sheet.add_image(Image(image), "E9")
    for col in "ABCDEF":
        sheet.column_dimensions[col].width = 19
    sheet.print_area = "A1:H24"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    hidden = workbook.create_sheet("Hidden")
    hidden["A1"] = "HIDDEN-NOT-READ"
    hidden.sheet_state = "hidden"
    source = io.BytesIO()
    workbook.save(source)
    # These are ORIGINAL file bytes: freeze B6's stored cache as 60 while D6
    # has no cache. A later LibreOffice recalculation is a separate observation.
    patched = io.BytesIO()
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(patched, "w") as archive:
        for item in original.infolist():
            raw = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(raw)
                root.find(".//m:c[@r='B6']/m:v", ns).text = "60"
                raw = ET.tostring(root)
            archive.writestr(item, raw)
    target.write_bytes(patched.getvalue())
    # Explicit author-defined facts, not a recount of converter/model output.
    oracle = {
        "source_sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
        "sheet": "Facts",
        "range": "A1:F6",
        "total_cells": 36,
        "non_empty_cells": 24,
        "empty_cells": 12,
        "non_empty_by_column": {"A": 6, "B": 5, "C": 4, "D": 5, "E": 2, "F": 2},
        "non_empty_coordinates": [
            "A1",
            "A2",
            "B2",
            "C2",
            "D2",
            "E2",
            "F2",
            "A3",
            "B3",
            "C3",
            "D3",
            "E3",
            "F3",
            "A4",
            "B4",
            "C4",
            "D4",
            "A5",
            "B5",
            "C5",
            "D5",
            "A6",
            "B6",
            "D6",
        ],
        "merge": "A1:C1",
        "revenue": [10, 20, 30],
        "growth": [0.125, 0.25, 0.375],
        "dates": ["2026-01-15", "2026-04-15", "2026-07-15"],
        "number_formats": {"B3:B5": '"$"#,##0.00', "C3:C5": "0.0%", "D3:D5": "yyyy-mm-dd"},
        "source_formulas": {
            "B6": {"formula": "=SUM(B3:B5)", "cache": 60},
            "D6": {"formula": "=SUM(B3:B5)*2", "cache": None},
        },
        "image": {"marker": "EMBEDDED-ORANGE-7C19", "background": "orange"},
        "chart": {
            "title": "Quarterly revenue",
            "categories": ["Q1", "Q2", "Q3"],
            "values": [10, 20, 30],
            "colors": ["blue", "orange", "green"],
        },
        "hidden_sheet_unread": "Hidden",
    }
    oracle_path.write_text(json.dumps(oracle, indent=2) + "\n")
    print(json.dumps({"source_sha256": oracle["source_sha256"], "oracle_frozen": True}))


if __name__ == "__main__":
    generate(Path(sys.argv[1]))
