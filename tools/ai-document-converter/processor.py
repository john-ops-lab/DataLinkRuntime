"""Single-process conversion implementation; never imports Control or a Provider."""

from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import re
import sys
import uuid
import zipfile
from datetime import date, datetime, time
from pathlib import Path

REVISION = "dlr-documents-v1-docling2.138-lo26.8"
MAX_SOURCE = 6 * 1024 * 1024
MAX_DERIVED = 64 * 1024 * 1024
MAX_PAGES = 200
MAX_PIXELS = 16_000_000
MAX_MEDIA = 10 * 1024 * 1024


class ConversionError(Exception):
    def __init__(self, code):
        self.code = code
        super().__init__(code)


def decode(value, limit=MAX_SOURCE):
    try:
        data = base64.b64decode(value, validate=True)
    except (TypeError, ValueError) as exc:
        raise ConversionError("ai_attachment_invalid") from exc
    if len(data) > limit:
        raise ConversionError("ai_attachment_too_large")
    return data


def base_result(source):
    return dict(
        status="success",
        source_sha256=hashlib.sha256(source).hexdigest(),
        representation_revision=REVISION,
        page_count=0,
        pages=[],
        sheets=[],
        coverage={},
        warnings=[],
        assets=[],
        data={},
    )


def archive_policy(source):
    warnings = []
    with zipfile.ZipFile(io.BytesIO(source)) as archive:
        members = archive.infolist()
        if len(members) > 4000 or sum(x.file_size for x in members) > 64 * 1024 * 1024:
            raise ConversionError("ai_attachment_unsafe_archive")
        for member in members:
            if (
                member.file_size > 16 * 1024 * 1024
                or member.file_size > max(1, member.compress_size) * 50
            ):
                raise ConversionError("ai_attachment_unsafe_archive")
            if "vbaproject" in member.filename.casefold():
                warnings.append("office_macros_not_rendered")
            if member.filename.endswith(".rels"):
                import xml.etree.ElementTree as ET

                root = ET.fromstring(archive.read(member))
                if any(x.attrib.get("TargetMode") == "External" for x in root):
                    warnings.append("office_external_relations_not_rendered")
    return sorted(set(warnings))


def json_value(value):
    return value.isoformat() if isinstance(value, (datetime, date, time)) else value


def workbook(source, extension, sheet=None, range_a1=None):
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    cells = []
    sheets = []
    warnings = []
    merges = []
    bounds = None
    if range_a1 is not None:
        if not re.fullmatch(
            r"[A-Za-z]{1,3}[1-9][0-9]{0,5}(:[A-Za-z]{1,3}[1-9][0-9]{0,5})?", range_a1
        ):
            raise ConversionError("ai_attachment_invalid")
        bounds = range_boundaries(range_a1)
        c1, r1, c2, r2 = bounds
        if c2 < c1 or r2 < r1 or (c2 - c1 + 1) * (r2 - r1 + 1) > 200 or c2 > 16384 or r2 > 100000:
            raise ConversionError("ai_attachment_range_exceeded")
    if extension == ".xlsx":
        from openpyxl import load_workbook

        warnings = archive_policy(source)
        book = load_workbook(io.BytesIO(source), data_only=False, keep_links=False)
        cached = load_workbook(io.BytesIO(source), data_only=True, keep_links=False)
        try:
            for tab in book:
                if tab.max_row > 100000 or tab.max_column > 16384:
                    raise ConversionError("ai_attachment_too_large")
                sheets.append(
                    dict(
                        name=tab.title,
                        hidden=tab.sheet_state != "visible",
                        rows=tab.max_row,
                        columns=tab.max_column,
                        print_area=str(tab.print_area),
                        merged_ranges=[str(x) for x in tab.merged_cells.ranges],
                        chart_count=len(tab._charts),
                    )
                )
            if sheet is not None:
                if sheet not in book.sheetnames:
                    raise ConversionError("ai_attachment_sheet_not_found")
                tab = book[sheet]
                values = cached[sheet]
                merges = [str(x) for x in tab.merged_cells.ranges]
                c1, r1, c2, r2 = bounds
                for row in tab.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                    for cell in row:
                        formula = cell.value if cell.data_type == "f" else None
                        cache = values[cell.coordinate].value
                        cells.append(
                            dict(
                                coordinate=cell.coordinate,
                                raw_value=json_value(cell.value),
                                data_type=cell.data_type,
                                number_format=cell.number_format,
                                formula=formula,
                                cached_value=json_value(cache),
                                cache_status="missing" if formula and cache is None else "stored",
                                value_origin="source_file",
                                formula_status="available" if formula else "not_formula",
                            )
                        )
        finally:
            book.close()
            cached.close()
    else:
        import xlrd

        book = xlrd.open_workbook(file_contents=source, formatting_info=True, on_demand=True)
        try:
            for tab in book.sheets():
                if tab.nrows > 100000 or tab.ncols > 16384:
                    raise ConversionError("ai_attachment_too_large")
                tab_merges = [
                    f"{get_column_letter(c1 + 1)}{r1 + 1}:{get_column_letter(c2)}{r2}"
                    for r1, r2, c1, c2 in tab.merged_cells
                ]
                sheets.append(
                    dict(
                        name=tab.name,
                        hidden=bool(tab.visibility),
                        rows=tab.nrows,
                        columns=tab.ncols,
                        print_area=None,
                        merged_ranges=tab_merges,
                        chart_count=None,
                    )
                )
            warnings.extend(
                [
                    "xls_formula_text_unavailable",
                    "xls_print_area_and_chart_metadata_unavailable",
                ]
            )
            if sheet is not None:
                if sheet not in book.sheet_names():
                    raise ConversionError("ai_attachment_sheet_not_found")
                tab = book.sheet_by_name(sheet)
                merges = next(x["merged_ranges"] for x in sheets if x["name"] == sheet)
                c1, r1, c2, r2 = bounds
                for row in range(r1 - 1, r2):
                    for column in range(c1 - 1, c2):
                        cell = (
                            tab.cell(row, column)
                            if row < tab.nrows and column < tab.ncols
                            else None
                        )
                        fmt = (
                            book.format_map.get(book.xf_list[cell.xf_index].format_key)
                            if cell is not None
                            else None
                        )
                        value = cell.value if cell is not None else None
                        cells.append(
                            dict(
                                coordinate=f"{get_column_letter(column + 1)}{row + 1}",
                                raw_value=value,
                                data_type=cell.ctype if cell is not None else xlrd.XL_CELL_EMPTY,
                                number_format=fmt.format_str if fmt else None,
                                formula=None,
                                cached_value=value,
                                cache_status="source_stored_value",
                                formula_status="unavailable_xlrd",
                                value_origin="source_file",
                            )
                        )
        finally:
            book.release_resources()
    if bounds is not None:
        selected = []
        for merged in merges:
            mc1, mr1, mc2, mr2 = range_boundaries(merged)
            if mc1 <= bounds[2] and mc2 >= bounds[0] and mr1 <= bounds[3] and mr2 >= bounds[1]:
                selected.append(merged)
                if len(selected) == 200:
                    warnings.append("sheet_metadata_partial")
                    break
        merges = selected
    return (
        sheets,
        dict(
            sheet=sheet,
            range=range_a1,
            cells=cells,
            merged_ranges=merges,
            value_origin="source_file",
            rendered_values_origin="libreoffice_recalculation",
        ),
        warnings,
    )


def png_asset(image):
    if image.width * image.height > MAX_PIXELS:
        raise ConversionError("ai_attachment_image_too_large")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    data = stream.getvalue()
    if len(data) > MAX_MEDIA:
        raise ConversionError("ai_attachment_too_large")
    return dict(
        asset_id=uuid.uuid4().hex,
        content_type="image/png",
        data_base64=base64.b64encode(data).decode(),
        sha256=hashlib.sha256(data).hexdigest(),
        size_bytes=len(data),
        width=image.width,
        height=image.height,
    )


def pdf_catalog(pdf, sheets):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(pdf), strict=False)
    if reader.is_encrypted:
        raise ConversionError("ai_attachment_encrypted")
    count = len(reader.pages)
    if count < 1 or count > MAX_PAGES:
        raise ConversionError("ai_attachment_page_limit")
    starts = []

    def outlines(items):
        for item in items:
            if isinstance(item, list):
                outlines(item)
            elif item.title in {s["name"] for s in sheets}:
                starts.append((reader.get_destination_page_number(item) + 1, item.title))

    if sheets:
        outlines(reader.outline)
    starts = sorted(set(starts))
    pages = []
    for index, page in enumerate(reader.pages, 1):
        mapped = [name for start, name in starts if start <= index]
        text = page.extract_text() or ""
        pages.append(
            dict(
                page=index,
                source_sheet=mapped[-1] if mapped else None,
                sheet_mapping="pdf_outline" if mapped else "unknown" if sheets else "original_pdf",
                text_layer=bool(text.strip()),
                status="available",
            )
        )
    for index, (start, name) in enumerate(starts):
        end = starts[index + 1][0] - 1 if index + 1 < len(starts) else count
        next(s for s in sheets if s["name"] == name)["rendered_page_start"] = start
        next(s for s in sheets if s["name"] == name)["rendered_page_end"] = end
    return reader, pages


def read_pdf_page(pdf, page_number, use_ocr, extract_structure=True):
    import pypdfium2
    from PIL import ImageStat
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(pdf))
    text = reader.pages[page_number - 1].extract_text() or ""
    pdfium = pypdfium2.PdfDocument(pdf)
    page = pdfium[page_number - 1]
    try:
        width, height = page.get_size()
        if int(width * 1.5) * int(height * 1.5) > MAX_PIXELS:
            raise ConversionError("ai_attachment_image_too_large")
        bitmap = page.render(scale=1.5)
        try:
            image = bitmap.to_pil().copy()
        finally:
            bitmap.close()
    finally:
        page.close()
        pdfium.close()
    asset = png_asset(image)
    blank = (
        min(ImageStat.Stat(image.convert("RGB")).mean) > 254.5
        and max(ImageStat.Stat(image.convert("RGB")).var) < 0.2
    )
    data = dict(
        page=page_number,
        text=text[:65536],
        tables=[],
        pictures=[],
        blank=blank,
        text_origin="pdf_text_layer",
        structure_status="skipped" if blank else "unavailable",
    )
    if blank:
        return asset, data, []
    if not extract_structure:
        data.update(text="", text_origin="not_requested", structure_status="not_requested")
        return asset, data, []
    from docling.datamodel.accelerator_options import (
        AcceleratorDevice,
        AcceleratorOptions,
    )
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions, RapidOcrOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    options = PdfPipelineOptions(
        artifacts_path=Path("/opt/app-root/src/.cache/docling/models"),
        do_ocr=use_ocr,
        generate_page_images=False,
        generate_picture_images=True,
        images_scale=1.5,
        accelerator_options=AcceleratorOptions(device=AcceleratorDevice.CPU, num_threads=2),
        ocr_options=RapidOcrOptions(backend="torch", lang=["ch"]),
    )
    converter = DocumentConverter(
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)}
    )
    from docling.datamodel.base_models import DocumentStream

    converted = converter.convert(
        DocumentStream(name="source.pdf", stream=io.BytesIO(pdf)),
        raises_on_error=False,
        page_range=(page_number, page_number),
        max_num_pages=MAX_PAGES,
        max_file_size=MAX_DERIVED,
    )
    doc = converted.document.export_to_dict()
    warnings = []
    if converted.status.value not in {"success", "partial_success"}:
        warnings.append("document_structure_unavailable")
    else:
        extracted = "\n".join(x.get("text", "") for x in doc.get("texts", []))
        data["text"] = extracted[:65536] or text[:65536]
        data["text_origin"] = "docling_ocr" if use_ocr else "docling_pdf_text"
        data["tables"] = [
            {"provenance": t.get("prov", []), "data": t.get("data", {})}
            for t in doc.get("tables", [])
        ][:20]
        data["pictures"] = [
            {"provenance": t.get("prov", []), "label": t.get("label")}
            for t in doc.get("pictures", [])
        ][:20]
        data["structure_status"] = (
            "partial"
            if data["tables"] or converted.status.value == "partial_success"
            else "success"
        )
        if data["tables"]:
            warnings.append("structure_partial")
    return asset, data, warnings


def dispatch(request):
    source = decode(request["data_base64"])
    result = base_result(source)
    extension = Path(request["filename"]).suffix.lower()
    operation = request["operation"]
    result["warnings"].extend(request.get("office_warnings", []))
    if extension not in {
        ".pdf",
        ".xlsx",
        ".xls",
        ".docx",
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
    }:
        raise ConversionError("ai_attachment_type_unsupported")
    if extension in {".xlsx", ".xls"}:
        result["sheets"], result["data"], warnings = workbook(
            source,
            extension,
            request.get("sheet") if operation == "read-sheet" else None,
            request.get("range_a1") if operation == "read-sheet" else None,
        )
        result["warnings"].extend(warnings)
    if operation == "read-sheet":
        if extension not in {".xlsx", ".xls"}:
            raise ConversionError("ai_attachment_type_unsupported")
        result["coverage"] = {
            "cells": "exact_source_range",
            "formula_execution": "never",
        }
        if any(x.get("cache_status") == "missing" for x in result["data"]["cells"]):
            result["warnings"].append("formula_cache_missing")
        return result
    if extension in {".png", ".jpg", ".jpeg", ".webp"}:
        from PIL import Image, ImageOps

        with Image.open(io.BytesIO(source)) as image:
            if image.width * image.height > MAX_PIXELS:
                raise ConversionError("ai_attachment_image_too_large")
            asset = png_asset(ImageOps.exif_transpose(image).convert("RGB"))
        result.update(
            page_count=1,
            pages=[dict(page=1, status="ready", image_asset_id=asset["asset_id"])],
            assets=[asset],
            coverage={
                "pages": "success",
                "read_pages": 1,
                "total_pages": 1,
                "structure": "not_applicable",
            },
        )
        return result
    if extension == ".pdf":
        pdf = source
    else:
        encoded = request.get("rendered_pdf_base64")
        if not encoded:
            result.update(
                status="partial_success",
                coverage={"pages": "unavailable", "cells": "source_metadata"},
            )
            result["warnings"].append("office_rendering_unavailable")
            return result
        pdf = decode(encoded, MAX_DERIVED)
        if operation == "prepare":
            result["rendered_pdf"] = encoded
    reader, result["pages"] = pdf_catalog(pdf, result["sheets"])
    result["page_count"] = len(result["pages"])
    number = int(request.get("page", 1))
    if number < 1 or number > result["page_count"]:
        raise ConversionError("ai_attachment_page_not_found")
    if request.get("view", "both") not in {"image", "text", "both", "pdf"}:
        raise ConversionError("ai_attachment_invalid")
    use_ocr = not result["pages"][number - 1]["text_layer"]
    if request.get("view") == "pdf":
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_page(reader.pages[number - 1])
        stream = io.BytesIO()
        writer.write(stream)
        content = stream.getvalue()
        if len(content) > MAX_MEDIA:
            raise ConversionError("ai_attachment_too_large")
        asset = {
            "asset_id": uuid.uuid4().hex,
            "content_type": "application/pdf",
            "data_base64": base64.b64encode(content).decode(),
            "sha256": hashlib.sha256(content).hexdigest(),
            "size_bytes": len(content),
            "page_count": 1,
        }
        result["assets"] = [asset]
        result["data"] = {
            "page": number,
            "text_origin": "original_pdf_page"
            if extension == ".pdf"
            else "libreoffice_rendered_pdf_page",
        }
        result["pages"][number - 1].update(status="ready", image_asset_id=asset["asset_id"])
        result["coverage"] = {
            "pages": "partial",
            "read_pages": 1,
            "total_pages": result["page_count"],
            "structure": "not_requested",
            "ocr": "skipped",
        }
        return result
    asset, data, warnings = read_pdf_page(
        pdf, number, use_ocr, extract_structure=request.get("view", "both") != "image"
    )
    result["assets"] = [asset] if request.get("view", "both") != "text" else []
    result["data"] = data
    result["warnings"].extend(warnings)
    result["pages"][number - 1].update(
        status="blank" if data["blank"] else "ready",
        structure_status=data["structure_status"],
    )
    if result["assets"]:
        result["pages"][number - 1]["image_asset_id"] = asset["asset_id"]
    result["coverage"] = {
        "pages": "partial",
        "read_pages": 1,
        "total_pages": result["page_count"],
        "structure": data["structure_status"],
        "ocr": "success"
        if data["text_origin"] == "docling_ocr"
        else "unavailable"
        if use_ocr and data["structure_status"] == "unavailable"
        else "skipped",
        "unread_pages": result["page_count"] - 1,
    }
    if warnings:
        result["status"] = "partial_success"
    if extension != ".pdf":
        result["warnings"].extend(["rendered_values_recalculated", "printed_area_only"])
    return result


if __name__ == "__main__":
    os.umask(0o077)
    from sandbox import restrict_child

    restrict_child()
    try:
        result = dispatch(json.loads(Path(sys.argv[1]).read_text()))
        body = json.dumps(result, ensure_ascii=False).encode()
        if len(body) > 96 * 1024 * 1024:
            raise ConversionError("ai_attachment_too_large")
        Path(sys.argv[2]).write_bytes(body)
    except ConversionError as error:
        Path(sys.argv[2]).write_text(json.dumps({"code": error.code}))
        sys.exit(2)
    except Exception:
        Path(sys.argv[2]).write_text(json.dumps({"code": "ai_attachment_parse_failed"}))
        sys.exit(2)
