"""Bounded source metadata, lazy PDF provenance, and process-supervisor contracts."""

from __future__ import annotations

import base64
import importlib.util
import io
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter

from test_xls_fixture import build_xls

TOOLS = Path(__file__).resolve().parents[2] / "tools" / "ai-document-converter"


def _module(name: str) -> Any:
    spec = importlib.util.spec_from_file_location(f"issue202_{name}", TOOLS / f"{name}.py")
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


processor = _module("processor")


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Summary"
    sheet["B2"] = 10
    sheet["B2"].number_format = "#,##0.00"
    sheet["C2"] = 0.125
    sheet["C2"].number_format = "0.0%"
    sheet["B3"] = "=SUM(B2:B2)"
    sheet.print_area = "A1:C5"
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_source_xlsx_cells_never_substitute_recalculated_values() -> None:
    source = _xlsx()
    sheets, data, warnings = processor.workbook(source, ".xlsx", "Sales", "B2:C3")
    assert sheets[1]["hidden"] is True
    assert sheets[0]["merged_ranges"] == ["A1:C1"]
    assert sheets[0]["print_area"] == "'Sales'!$A$1:$C$5"
    rows = {cell["coordinate"]: cell for cell in data["cells"]}
    assert rows["B2"]["raw_value"] == 10
    assert rows["C2"]["number_format"] == "0.0%"
    assert rows["B3"]["formula"] == "=SUM(B2:B2)"
    assert rows["B3"]["cached_value"] is None
    assert rows["B3"]["cache_status"] == "missing"
    assert warnings == []
    assert data["merged_ranges"] == []


@pytest.mark.parametrize("range_a1", ["A1:A201", "A0", "B2:A1", "A1:ZZZ2", "../A1"])
def test_range_is_bounded_before_iteration(range_a1: str) -> None:
    with pytest.raises(processor.ConversionError):
        processor.workbook(_xlsx(), ".xlsx", "Sales", range_a1)


def test_true_biff_uses_stored_values_and_marks_formula_text_unknown() -> None:
    source = build_xls()
    assert source.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    sheets, data, warnings = processor.workbook(source, ".xls")
    name = sheets[0]["name"]
    _, data, _ = processor.workbook(source, ".xls", name, "A1:B2")
    assert len(data["cells"]) == 4
    assert all(x["formula_status"] == "unavailable_xlrd" for x in data["cells"])
    assert "xls_formula_text_unavailable" in warnings


def test_outline_sheet_mapping_uses_real_page_destinations() -> None:
    writer = PdfWriter()
    for _ in range(4):
        writer.add_blank_page(width=595, height=842)
    writer.add_outline_item("Summary", 0)
    writer.add_outline_item("Details", 1)
    stream = io.BytesIO()
    writer.write(stream)
    sheets = [{"name": "Summary"}, {"name": "Details"}]
    _, pages = processor.pdf_catalog(stream.getvalue(), sheets)
    assert [x["page"] for x in pages] == [1, 2, 3, 4]
    assert [x["source_sheet"] for x in pages] == ["Summary", "Details", "Details", "Details"]
    assert sheets[1]["rendered_page_start"] == 2
    assert sheets[1]["rendered_page_end"] == 4


def test_native_single_page_export_preserves_original_ordinal() -> None:
    writer = PdfWriter()
    for width in [100, 200, 300]:
        writer.add_blank_page(width=width, height=400)
    stream = io.BytesIO()
    writer.write(stream)
    result = processor.dispatch(
        {
            "operation": "read-page",
            "filename": "original.pdf",
            "data_base64": base64.b64encode(stream.getvalue()).decode(),
            "page": 3,
            "view": "pdf",
        }
    )
    asset = result["assets"][0]
    extracted = PdfReader(io.BytesIO(base64.b64decode(asset["data_base64"])))
    assert len(extracted.pages) == 1
    assert extracted.pages[0].mediabox.width == 300
    assert result["data"]["page"] == 3
    assert result["page_count"] == 3
    assert asset["page_count"] == 1


def test_cancel_before_registration_remains_terminal(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.syspath_prepend(str(TOOLS))
    # server's processor import intentionally has no document-library side effects.
    module = _module("server")
    monkeypatch.setattr(module, "ROLE", "office")
    supervisor = module.Supervisor()
    assert supervisor.cancel("cancel-first")["terminated"] is True
    with pytest.raises(module.ConversionError):
        supervisor.run(
            {
                "job_id": "cancel-first",
                "filename": "source.pdf",
                "data_base64": base64.b64encode(b"source").decode(),
            },
            "prepare",
        )
    assert supervisor.jobs == {}
